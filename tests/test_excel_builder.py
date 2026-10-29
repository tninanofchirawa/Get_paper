"""
tests/test_excel_builder.py - Verifies Section 5 standardized Excel output layout.
"""

import os
import unittest
import openpyxl

from config import cfg
from excel_builder import StandardizedExcelBuilder
from harvester.models import Paper
from analyzer import PaperAnalyzer


class TestStandardizedExcelBuilder(unittest.TestCase):
    def setUp(self):
        self.test_output = os.path.join(cfg.OUTPUT_DIR, "test_standardized_output.xlsx")
        if os.path.exists(self.test_output):
            os.remove(self.test_output)

    def tearDown(self):
        if os.path.exists(self.test_output):
            try:
                os.remove(self.test_output)
            except Exception:
                pass

    def test_excel_metadata_and_table_structure(self):
        papers = [
            Paper(
                title="Fuzzy Logic Foundations",
                authors="Lotfi A. Zadeh",
                year="1988",
                venue="Computer",
                doi="10.1109/2.53",
                url="https://doi.org/10.1109/2.53",
                pdf_link="https://ieeexplore.ieee.org/stamp/stamp.jsp?tp=&arnumber=53",
                source="Crossref",
                scopus_indexed=True,
                wos_indexed=True,
                ieee_indexed=True,
                open_access=False,
                verification_source="Scopus API, WoS API, IEEE API"
            )
        ]
        analyzed = [PaperAnalyzer.analyze_paper(papers[0], query="fuzzy logic", material_idx=1)]

        run_config = {
            "run_name": "Test_Run_20261028",
            "run_type": "Automated Scheduled",
            "location": "HITS Campus",
            "main_topics": ["Fuzzy Set Theory"],
            "subtopics": ["Alpha Cuts", "Extension Principle", "Interval Numbers"],
            "harvest_quota": 250,
            "year_range": "2021-2026",
            "doc_filter": "Journal Articles Only",
            "inst_token_flag": True,
            "target_repos": ["Scopus", "Web of Science", "IEEE Xplore"],
            "quota_strictness": "Strict"
        }
        metrics = {
            "raw_scraped": 300,
            "duplicates_removed": 50,
            "verified_retained": 1
        }

        excel_path = StandardizedExcelBuilder.generate_workbook(
            papers=analyzed,
            run_config=run_config,
            metrics=metrics,
            output_filepath=self.test_output
        )

        self.assertTrue(os.path.exists(excel_path))

        # Inspect generated workbook
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb.active

        # 1. Assert Column A (Rows 1 to 10) Primary Metadata Labels
        self.assertEqual(ws["A1"].value, "Main Topic")
        self.assertEqual(ws["A2"].value, "Subtopic 1")
        self.assertEqual(ws["A3"].value, "Subtopic 2")
        self.assertEqual(ws["A4"].value, "Subtopic 3")
        self.assertEqual(ws["A7"].value, "Survey Question 1 (Harvest Quota & Year Range)")
        self.assertEqual(ws["A8"].value, "Survey Question 2 (Document Classification Filter)")
        self.assertEqual(ws["A9"].value, "Survey Question 3 (Institutional Token & Target Repositories)")
        self.assertEqual(ws["A10"].value, "Survey Question 4 (Quota Strictness & Execution Mode)")

        # 2. Assert Column B (Rows 1 to 10) Values
        self.assertEqual(ws["B1"].value, "Fuzzy Set Theory")
        self.assertEqual(ws["B2"].value, "Alpha Cuts")
        self.assertEqual(ws["B3"].value, "Extension Principle")
        self.assertIn("Target: 250", str(ws["B7"].value))

        # 3. Assert Column C (Rows 1 to 10) Empty Spacer Column
        for r in range(1, 11):
            self.assertIsNone(ws[f"C{r}"].value)

        # 4. Assert Column D & E (Rows 1 to 10) Run Configuration
        self.assertEqual(ws["D1"].value, "Run Name")
        self.assertEqual(ws["E1"].value, "Test_Run_20261028")
        self.assertEqual(ws["D2"].value, "Run Type")
        self.assertEqual(ws["E2"].value, "Automated Scheduled")
        self.assertEqual(ws["D4"].value, "Execution Location")
        self.assertEqual(ws["E4"].value, "HITS Campus")
        self.assertEqual(ws["D8"].value, "Target Quota")
        self.assertEqual(ws["E8"].value, "250")
        self.assertEqual(ws["D9"].value, "Quota Reached / Verified Retained")
        self.assertEqual(ws["E9"].value, "1")

        # 5. Assert Row 11 Header Row
        headers_in_row_11 = [ws.cell(11, col).value for col in range(1, len(StandardizedExcelBuilder.TABLE_HEADERS) + 1)]
        self.assertIn("Title", headers_in_row_11)
        self.assertIn("Authors", headers_in_row_11)
        self.assertIn("Publication Year", headers_in_row_11)
        self.assertIn("Journal/Source", headers_in_row_11)
        self.assertIn("DOI", headers_in_row_11)
        self.assertIn("URL", headers_in_row_11)
        self.assertIn("Scopus Indexed (Y/N)", headers_in_row_11)
        self.assertIn("WoS Indexed (Y/N)", headers_in_row_11)
        self.assertIn("IEEE Indexed (Y/N)", headers_in_row_11)
        self.assertIn("OpenAccess (Y/N)", headers_in_row_11)
        self.assertIn("Verification Source", headers_in_row_11)

        # 6. Assert Row 12 First Paper Record
        self.assertEqual(ws["A12"].value, "M-0001")
        self.assertEqual(ws["B12"].value, "Fuzzy Logic Foundations")
        self.assertEqual(ws["C12"].value, "Lotfi A. Zadeh")


if __name__ == "__main__":
    unittest.main()
